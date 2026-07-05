#!/usr/bin/env python3
"""Build a student Word proof pack from organized comprehensive-evaluation images."""

from __future__ import annotations

import argparse
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from openpyxl import load_workbook
from PIL import Image


TEMPLATE_RELATIVE_PATH = Path("综测模板") / "张三（提交示例！）" / "张三综测证明材料.docx"
CATEGORY_ORDER = ("身心素养", "文艺素养", "劳动素养", "创新素养")
SUBCATEGORY_ORDER = ("基础性评价", "成果性评价")
IMAGE_EXTS = {".jpg", ".jpeg", ".png"}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Create <student>综测证明材料.docx from organized proof-image folders."
    )
    parser.add_argument("student_folder", type=Path, help="学生提交文件夹")
    parser.add_argument("--template", type=Path, help="Word 证明材料模板")
    parser.add_argument("--out", type=Path, help="输出 DOCX；默认写入学生文件夹")
    return parser.parse_args()


def default_template_candidates(student_folder: Path) -> list[Path]:
    script_path = Path(__file__).resolve()
    repo_root = script_path.parents[3] if len(script_path.parents) > 3 else script_path.parent
    return [
        student_folder.parent / TEMPLATE_RELATIVE_PATH,
        Path.cwd() / TEMPLATE_RELATIVE_PATH,
        Path.cwd() / "project-files" / TEMPLATE_RELATIVE_PATH,
        repo_root / "project-files" / TEMPLATE_RELATIVE_PATH,
    ]


def find_template(student_folder: Path, template_arg: Path | None) -> Path:
    if template_arg is not None:
        return template_arg.resolve()
    for candidate in default_template_candidates(student_folder):
        if candidate.is_file():
            return candidate.resolve()
    tried = "\n".join(str(path) for path in default_template_candidates(student_folder))
    raise SystemExit(f"模板不存在。请用 --template 指定；已尝试:\n{tried}")


def read_student_info(student_folder: Path) -> tuple[str, str, str]:
    fallback_name = student_folder.name
    workbooks = sorted(student_folder.glob("*综测统计.xlsx"))
    if not workbooks:
        return "", "", fallback_name

    wb = load_workbook(workbooks[0], data_only=True)
    ws = wb.active
    student_class = str(ws["B7"].value or "")
    student_id = str(ws["C7"].value or "")
    student_name = str(ws["D7"].value or fallback_name)
    return student_class, student_id, student_name


def collect_materials(student_folder: Path) -> dict[str, dict[str, list[Path]]]:
    materials: dict[str, dict[str, list[Path]]] = {
        category: {subcategory: [] for subcategory in SUBCATEGORY_ORDER}
        for category in CATEGORY_ORDER
    }

    for category in CATEGORY_ORDER:
        category_dirs = sorted(
            path for path in student_folder.iterdir() if path.is_dir() and category in path.name
        )
        for category_dir in category_dirs:
            for subcategory in SUBCATEGORY_ORDER:
                sub_dirs = sorted(
                    path for path in category_dir.iterdir() if path.is_dir() and subcategory in path.name
                )
                for sub_dir in sub_dirs:
                    images = sorted(
                        path
                        for path in sub_dir.rglob("*")
                        if path.is_file() and path.suffix.lower() in IMAGE_EXTS
                    )
                    materials[category][subcategory].extend(images)

    return materials


def clear_document_body(doc: Document) -> None:
    body = doc._body._element
    sect_pr = None
    for child in list(body):
        if child.tag.endswith("}sectPr"):
            sect_pr = deepcopy(child)
        body.remove(child)
    if sect_pr is not None:
        body.append(sect_pr)


def set_run_font(run, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = "SimSun"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_paragraph(
    doc: Document,
    text: str,
    *,
    size: float = 12,
    bold: bool = False,
    align: int | None = None,
    space_after: float = 4,
    page_break_before: bool = False,
):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.page_break_before = page_break_before
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return paragraph


def refresh_header(doc: Document, header_text: str) -> None:
    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            paragraph.text = ""
        header = section.header.paragraphs[0] if section.header.paragraphs else section.header.add_paragraph()
        header.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = header.add_run(header_text)
        set_run_font(run, size=9)
        for paragraph in section.footer.paragraphs:
            paragraph.text = ""


def add_image_block(
    doc: Document,
    *,
    category: str,
    subcategory: str,
    image_path: Path,
    first_page: bool,
    student_name: str,
    max_img_width: int,
    max_img_height: int,
) -> None:
    if first_page:
        add_paragraph(
            doc,
            f"{student_name}综测证明材料",
            size=15,
            bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER,
            space_after=8,
        )

    add_paragraph(
        doc,
        f"【{category}】",
        size=13,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        page_break_before=not first_page,
    )
    add_paragraph(
        doc,
        f"{subcategory}证明材料",
        size=11,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.LEFT,
    )
    add_paragraph(
        doc,
        image_path.stem,
        size=10.5,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    with Image.open(image_path) as image:
        width_px, height_px = image.size
    scale = min(max_img_width / width_px, max_img_height / height_px)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.add_run().add_picture(
        str(image_path),
        width=int(width_px * scale),
        height=int(height_px * scale),
    )


def add_empty_category(doc: Document, category: str, *, first_page: bool) -> None:
    add_paragraph(
        doc,
        f"【{category}】",
        size=13,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        space_after=8,
        page_break_before=not first_page,
    )
    add_paragraph(doc, "本栏目暂无图片证明材料。", size=11, align=WD_ALIGN_PARAGRAPH.LEFT)


def build_proof_pack(student_folder: Path, template: Path, out: Path) -> None:
    student_class, student_id, student_name = read_student_info(student_folder)
    header_text = f"{student_class} {student_id} {student_name} 综测证明材料".strip()
    materials = collect_materials(student_folder)

    doc = Document(template)
    section = doc.sections[0]
    usable_width = section.page_width - section.left_margin - section.right_margin
    max_img_width = int(usable_width * 0.96)
    max_img_height = int(Inches(6.75))

    clear_document_body(doc)
    refresh_header(doc, header_text)

    first_page = True
    for category in CATEGORY_ORDER:
        category_images = [
            (subcategory, image)
            for subcategory in SUBCATEGORY_ORDER
            for image in materials[category][subcategory]
        ]
        if not category_images:
            add_empty_category(doc, category, first_page=first_page)
            first_page = False
            continue

        for subcategory, image_path in category_images:
            add_image_block(
                doc,
                category=category,
                subcategory=subcategory,
                image_path=image_path,
                first_page=first_page,
                student_name=student_name,
                max_img_width=max_img_width,
                max_img_height=max_img_height,
            )
            first_page = False

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)

    for category in CATEGORY_ORDER:
        count = sum(len(materials[category][subcategory]) for subcategory in SUBCATEGORY_ORDER)
        print(f"{category}: {count} 张")
    print(out)


def main() -> None:
    args = parse_args()
    student_folder = args.student_folder.resolve()
    template = find_template(student_folder, args.template)
    out = args.out.resolve() if args.out else student_folder / f"{student_folder.name}综测证明材料.docx"

    if not student_folder.is_dir():
        raise SystemExit(f"学生文件夹不存在: {student_folder}")
    if not template.is_file():
        raise SystemExit(f"模板不存在: {template}")

    build_proof_pack(student_folder, template, out)


if __name__ == "__main__":
    main()
