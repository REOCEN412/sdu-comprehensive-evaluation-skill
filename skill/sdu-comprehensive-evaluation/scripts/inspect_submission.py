#!/usr/bin/env python3
"""Inspect SDU comprehensive evaluation submission folders.

This script is intentionally conservative: it reports likely problems but does
not move, rename, delete, or rewrite user files.
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

try:
    import openpyxl
    from openpyxl.utils.cell import range_boundaries
except Exception as exc:  # pragma: no cover - environment message
    openpyxl = None
    OPENPYXL_ERROR = exc
else:
    OPENPYXL_ERROR = None


CATEGORY_NAMES = ("身心素养", "文艺素养", "劳动素养", "创新素养")
SUBCATEGORY_NAMES = ("基础性评价", "成果性评价")
PROOF_SUFFIXES = {
    ".pdf",
    ".doc",
    ".docx",
    ".jpg",
    ".jpeg",
    ".png",
    ".heic",
    ".bmp",
    ".tif",
    ".tiff",
}

SCORE_COLUMNS = {
    "G": ("身心基础", 9),
    "I": ("身心成果", 6),
    "K": ("文艺基础", 9),
    "M": ("文艺成果", 6),
    "O": ("劳动基础", 15),
    "Q": ("劳动成果", 10),
    "S": ("创新基础", 5),
    "U": ("创新成果", 40),
    "V": ("总分", 100),
}


def is_student_dir(path: Path) -> bool:
    if not path.is_dir() or path.name.startswith("."):
        return False
    if path.name in {".codex", ".git", "综测模板"}:
        return False
    if any(category in path.name for category in CATEGORY_NAMES):
        return False
    return True


def find_named_dirs(path: Path, keyword: str) -> list[Path]:
    return [p for p in sorted(path.iterdir()) if p.is_dir() and keyword in p.name]


def has_category_dirs(path: Path) -> bool:
    return any(find_named_dirs(path, category) for category in CATEGORY_NAMES)


def parse_score(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    return float(match.group(0)) if match else None


def score_like_filename(path: Path) -> bool:
    stem = path.stem
    return bool(re.search(r"(\+?\d+(?:\.\d+)?\s*分|无|0\s*分)", stem))


def score_like_name(name: str) -> bool:
    return bool(re.search(r"(\d+(?:\.\d+)?\s*分|无|0\s*分)", name))


def merged_range_intersects(range_text: str, min_col: int, min_row: int, max_col: int, max_row: int) -> bool:
    left, top, right, bottom = range_boundaries(range_text)
    return not (right < min_col or left > max_col or bottom < min_row or top > max_row)


def is_allowed_root_proof(path: Path) -> bool:
    return path.suffix.lower() == ".docx" and "综测证明材料" in path.stem


def inspect_excel(path: Path) -> list[str]:
    issues: list[str] = []
    if openpyxl is None:
        return [f"无法读取 Excel：缺少 openpyxl ({OPENPYXL_ERROR})"]

    try:
        wb = openpyxl.load_workbook(path, data_only=False)
    except Exception as exc:
        return [f"无法打开 Excel：{path.name}: {exc}"]

    ws = wb.active
    title = str(ws["A1"].value or "")
    if "本科生素质能力评价成绩" not in title:
        issues.append(f"Excel 标题疑似不匹配：{path.name}")

    sample_merge_ranges = [
        str(rng)
        for rng in ws.merged_cells.ranges
        if merged_range_intersects(str(rng), 1, 7, 5, 8)
    ]
    if sample_merge_ranges:
        issues.append(
            "Excel 仍保留示例合并区域 A7:E8，正式学生行可能隐藏基本信息："
            + ", ".join(sample_merge_ranges)
        )

    basic_cells = {
        "A7": "序号",
        "B7": "专业班级",
        "C7": "学号",
        "D7": "姓名",
        "E7": "德育考核等级",
    }
    for cell, label in basic_cells.items():
        value = ws[cell].value
        if value is None or str(value).strip() in {"", "示例"}:
            issues.append(f"Excel {cell} {label} 缺失或仍为示例内容")

    for row in range(8, min(ws.max_row, 20) + 1):
        row_text = " ".join(str(ws.cell(row, col).value or "") for col in range(1, min(ws.max_column, 23) + 1))
        if "注意事项" in row_text:
            issues.append(f"Excel 第 {row} 行仍包含模板注意事项，最终提交版建议删除")
        if re.search(r"(^|\s)示例($|\s)", row_text):
            issues.append(f"Excel 第 {row} 行仍包含模板示例内容，最终提交版建议删除")

    for col, (label, cap) in SCORE_COLUMNS.items():
        value = parse_score(ws[f"{col}7"].value)
        if value is None:
            issues.append(f"Excel {col}7 {label} 未填写或无法识别分数")
        elif value < 0:
            issues.append(f"Excel {col}7 {label} 为负数：{value:g}")
        elif value > cap:
            issues.append(f"Excel {col}7 {label} 超上限 {cap}：{value:g}")

    reason_cells = {
        "F7": "身心基础理由",
        "H7": "身心成果理由",
        "J7": "文艺基础理由",
        "L7": "文艺成果理由",
        "N7": "劳动基础理由",
        "P7": "劳动成果理由",
        "R7": "创新基础理由",
        "T7": "创新成果理由",
    }
    for cell, label in reason_cells.items():
        score_col = chr(ord(cell[0]) + 1)
        score = parse_score(ws[f"{score_col}7"].value)
        reason = str(ws[cell].value or "").strip()
        if score and score > 0 and not reason:
            issues.append(f"Excel {cell} {label} 有分数但缺少理由")
        if reason and score == 0 and "+0" not in reason and "无" not in reason:
            issues.append(f"Excel {cell} {label} 有理由但对应分数为 0，建议写明 +0分 原因")

    return issues


def inspect_student_dir(path: Path) -> list[str]:
    issues: list[str] = []
    xlsx_files = sorted(path.glob("*.xlsx"))
    if not xlsx_files:
        issues.append("缺少综测统计 Excel")
    elif len(xlsx_files) > 1:
        issues.append("根目录存在多个 Excel，请确认哪一个是正式综测统计表")
    else:
        issues.extend(inspect_excel(xlsx_files[0]))

    for category in CATEGORY_NAMES:
        matches = find_named_dirs(path, category)
        if not matches:
            issues.append(f"缺少类别文件夹：{category}")
            continue
        if len(matches) > 1:
            issues.append(f"{category} 存在多个匹配文件夹，请确认是否重复：{', '.join(p.name for p in matches)}")
        category_dir = matches[0]
        if not score_like_name(category_dir.name):
            issues.append(f"{category} 类别文件夹名缺少总分：{category_dir.name}")
        subdir_names = [p.name for p in category_dir.iterdir() if p.is_dir()]
        for subcategory in SUBCATEGORY_NAMES:
            matching_subdirs = [name for name in subdir_names if subcategory in name]
            if not matching_subdirs:
                issues.append(f"{category} 缺少{subcategory}证明材料文件夹")
                continue
            for subdir_name in matching_subdirs:
                if "证明材料" not in subdir_name:
                    issues.append(f"{category} {subcategory}文件夹名建议包含“证明材料”：{subdir_name}")
                if not score_like_name(subdir_name):
                    issues.append(f"{category} {subcategory}文件夹名缺少分值：{subdir_name}")
        proof_files = [p for p in category_dir.rglob("*") if p.is_file() and p.suffix.lower() in PROOF_SUFFIXES]
        if proof_files and not any(any(name in subdir for name in SUBCATEGORY_NAMES) for subdir in subdir_names):
            issues.append(f"{category} 有证明材料但未区分 基础性评价/成果性评价")
        for proof in proof_files:
            if not score_like_filename(proof):
                issues.append(f"证明文件名缺少分值或无/0分说明：{proof.relative_to(path)}")
            if category == "创新素养":
                stem = proof.stem
                if not re.search(r"A\*?|A-?\*?|B\*?|C|D|E|类", stem, re.I):
                    issues.append(f"创新证明文件名建议写明赛事类别：{proof.relative_to(path)}")

    root_proofs = [
        p
        for p in path.iterdir()
        if p.is_file() and p.suffix.lower() in PROOF_SUFFIXES and p.suffix.lower() != ".xlsx"
    ]
    for proof in root_proofs:
        if is_allowed_root_proof(proof):
            continue
        issues.append(f"证明材料在学生根目录，建议归入对应类别：{proof.name}")

    return issues


def main() -> int:
    parser = argparse.ArgumentParser(description="Inspect SDU comprehensive evaluation submissions.")
    parser.add_argument("path", nargs="?", default=".", help="Student folder or parent folder to inspect")
    args = parser.parse_args()

    target = Path(args.path).expanduser().resolve()
    if not target.exists():
        print(f"路径不存在：{target}", file=sys.stderr)
        return 2

    student_dirs = [target]
    if target.is_dir() and not has_category_dirs(target):
        children = [p for p in sorted(target.iterdir()) if is_student_dir(p)]
        if children:
            student_dirs = children

    print(f"# 综测材料检查报告\n\n检查路径：`{target}`\n")
    total_issues = 0
    for student_dir in student_dirs:
        issues = inspect_student_dir(student_dir)
        total_issues += len(issues)
        print(f"## {student_dir.name}")
        if issues:
            for issue in issues:
                print(f"- {issue}")
        else:
            print("- 未发现明显结构/表格问题")
        print()

    print(f"合计问题数：{total_issues}")
    return 1 if total_issues else 0


if __name__ == "__main__":
    raise SystemExit(main())
